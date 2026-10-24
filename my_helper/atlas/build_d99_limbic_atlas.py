"""Extract the approved D99 subregions and merged masks without interpolation."""

from pathlib import Path
import csv
import json
import re
import tempfile

import nibabel as nib
import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SPEC = Path(__file__).with_name("specs") / "d99_limbic_regions.json"


def save_mask(mask, affine, reference, path):
    header = reference.header.copy()
    header.set_data_dtype(np.uint8)
    header.set_slope_inter(1, 0)
    header.set_intent("none")
    header["cal_min"], header["cal_max"] = 0, 1
    image = nib.Nifti1Image(mask.astype(np.uint8), affine, header)
    image.set_qform(affine, int(reference.header["qform_code"]))
    image.set_sform(affine, int(reference.header["sform_code"]))
    nib.save(image, path)


def main():
    spec = json.loads(SPEC.read_text())
    output = ROOT / spec["output_dir"]
    if output.exists():
        raise FileExistsError(f"Refusing to replace an existing atlas: {output}")
    reference = nib.load(ROOT / spec["reference_image"])
    label_image = nib.load(ROOT / spec["label_image"])
    if label_image.shape != reference.shape or not np.allclose(
        label_image.affine, reference.affine, atol=1e-6
    ):
        raise ValueError("D99 labels must already match the reference grid.")
    labels = np.asarray(label_image.dataobj)
    text_labels = {}
    for line_number, line in enumerate((ROOT / spec["label_text"]).read_text().splitlines(), 1):
        if line.strip():
            number, name = line.split(maxsplit=1)
            text_labels[int(number)] = (name, line_number)
    workbook = load_workbook(ROOT / spec["label_workbook"], read_only=True, data_only=True)
    sheet = workbook[spec["label_sheet"]]
    requested = {
        number for roi in spec["rois"] if roi["kind"] == "label"
        for number in roi["labels"].values()
    }
    metadata = {}
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row[1] in ("Left", "Right") and row[0] is not None:
            number = int(row[0])
            if number in requested:
                if number in metadata:
                    raise ValueError(f"Duplicate requested workbook label: {number}")
                metadata[number] = {
                    "side": row[1], "abbreviation": str(row[2]).strip(),
                    "full_name": str(row[3]).strip(), "workbook_row": row_number,
                }
    workbook.close()
    for roi in spec["rois"]:
        if roi["kind"] == "label":
            for side, number in roi["labels"].items():
                item = metadata[number]
                expected_side = "Left" if side == "lh" else "Right"
                if item["side"] != expected_side or item["abbreviation"] != roi["abbreviation"]:
                    raise ValueError(f"Workbook label mismatch: {roi['name']} {side}")
                if not text_labels[number][0].endswith("_" + expected_side):
                    raise ValueError(f"Text-label laterality mismatch: {number}")
                normalized_name = re.sub("[^a-z0-9]", "", item["full_name"].lower())
                normalized_text = re.sub("[^a-z0-9]", "", text_labels[number][0].lower())
                if not normalized_text.startswith(normalized_name):
                    raise ValueError(f"Workbook/text anatomical-name mismatch: {number}")
    source_files = [ROOT / spec[key] for key in (
        "reference_image", "label_image", "label_text", "label_workbook"
    )]
    source_files.extend(ROOT / roi["sources"][side] for roi in spec["rois"]
                        if roi["kind"] == "derived_mask" for side in ("lh", "rh"))
    preservation = {
        str(path): {"size": path.stat().st_size, "mtime_ns": path.stat().st_mtime_ns}
        for path in source_files
    }
    output.mkdir()
    work = Path(tempfile.mkdtemp(prefix="d99-limbic-index-", dir="/private/tmp"))
    records = []
    voxel_volume = float(abs(np.linalg.det(reference.affine[:3, :3])))
    for side in ("lh", "rh"):
        (output / side).mkdir()
        (work / side).mkdir()
        for roi in spec["rois"]:
            if roi["kind"] == "label":
                numbers = [roi["labels"][side]]
                mask = labels == numbers[0]
                full_name = metadata[numbers[0]]["full_name"]
                source_path = str(ROOT / spec["label_image"])
            elif roi["kind"] == "union":
                children = [r for r in spec["rois"] if r["name"] in roi["members"]]
                numbers = [r["labels"][side] for r in children]
                mask = np.isin(labels, numbers)
                full_name = "Approved merged " + roi["name"]
                source_path = str(ROOT / spec["label_image"])
            else:
                source_path = str(ROOT / roi["sources"][side])
                original = nib.load(source_path)
                data = np.asarray(original.dataobj)
                if not np.isin(data, [0, 1]).all():
                    raise ValueError(f"ALIC source is not binary: {source_path}")
                lattice = np.linalg.inv(reference.affine) @ original.affine
                offset = np.rint(lattice[:3, 3]).astype(int)
                if not np.allclose(lattice[:3, :3], np.eye(3), atol=1e-6) or not np.allclose(
                    lattice[:3, 3], offset, atol=1e-6
                ):
                    raise ValueError("ALIC source must lie on the same integer D99 lattice.")
                mask = np.zeros(reference.shape, dtype=bool)
                slices = tuple(slice(int(start), int(start + count))
                               for start, count in zip(offset, original.shape))
                mask[slices] = data.astype(bool)
                numbers = [roi["support_labels"][side]]
                if not np.all(labels[mask] == numbers[0]):
                    raise ValueError("Existing ALIC extends beyond its source white-matter label.")
                full_name = "Anterior limb of internal capsule (existing derived delineation)"
            coordinates = np.array(np.where(mask)).T
            if len(coordinates) == 0:
                raise ValueError(f"Requested ROI is empty: {roi['name']} {side}")
            world = nib.affines.apply_affine(reference.affine, coordinates)
            wrong_side = world[:, 0] > 1e-6 if side == "lh" else world[:, 0] < -1e-6
            if wrong_side.any():
                raise ValueError(f"ROI crosses into the wrong hemisphere: {roi['name']} {side}")
            path = output / side / (roi["name"] + ".nii.gz")
            save_mask(mask, reference.affine, reference, path)
            start = np.maximum(coordinates.min(0) - 2, 0)
            stop = np.minimum(coordinates.max(0) + 3, reference.shape)
            crop_slices = tuple(slice(int(lo), int(hi)) for lo, hi in zip(start, stop))
            crop_affine = reference.affine.copy()
            crop_affine[:3, 3] = nib.affines.apply_affine(reference.affine, start)
            crop_path = work / side / (roi["name"] + ".nii")
            save_mask(mask[crop_slices], crop_affine, reference, crop_path)
            records.append({
                "roi_name": roi["name"], "side": side, "family": roi["family"],
                "kind": roi["kind"], "parent": roi["parent"], "source_file": source_path,
                "source_labels": numbers, "abbreviation": roi.get("abbreviation", ""),
                "full_name": full_name, "members": roi.get("members", []),
                "workbook_rows": [metadata[n]["workbook_row"] for n in numbers if n in metadata],
                "text_rows": [text_labels[n][1] for n in numbers],
                "voxel_count": len(coordinates), "volume_mm3": len(coordinates) * voxel_volume,
                "centroid_mm": world.mean(0).tolist(), "wrong_hemisphere_voxels": int(wrong_side.sum()),
                "output_file": str(path), "cache_crop": str(crop_path),
            })
            print(f"Wrote {side}/{path.name}: {len(coordinates)} voxels", flush=True)
    payload = {
        "atlas_name": spec["atlas_name"], "space": spec["space"],
        "reference_image": str(ROOT / spec["reference_image"]),
        "label_sheet": spec["label_sheet"], "spec_path": str(SPEC),
        "output_dir": str(output), "index_work_dir": str(work),
        "parent_child_overlap": "Intentional; do not add parent and child measurements.",
        "ALIC_provenance": "Existing derived D99 subset; original boundary-authoring provenance unavailable.",
        "rois": records,
    }
    (output / "roi_manifest.json").write_text(json.dumps(payload, indent=2) + "\n")
    with (output / "roi_manifest.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=[key for key in records[0] if key != "cache_crop"])
        writer.writeheader()
        for record in records:
            writer.writerow({key: json.dumps(value) if isinstance(value, list) else value
                             for key, value in record.items() if key != "cache_crop"})
    (output / "source_inventory.json").write_text(json.dumps(preservation, indent=2) + "\n")
    (output / "README.md").write_text(Path(__file__).with_name("d99_limbic_atlas.md").read_text())
    print(f"Created {len(records)} masks in {output}", flush=True)


if __name__ == "__main__":
    main()
